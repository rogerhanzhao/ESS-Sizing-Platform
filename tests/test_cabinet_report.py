import copy
import io
import re
import zipfile

import pytest
from docx import Document

from tests.test_cabinet_report_data import saved_cabinet
from calb_diagrams.cabinet_sld import render_cabinet_sld
from calb_diagrams.cabinet_arrangement import render_cabinet_arrangement
from calb_sizing_tool.services.cabinet_report_data_service import build_cabinet_report_data
from calb_sizing_tool.reporting.cabinet_report import export_cabinet_report
from calb_sizing_tool.reporting.brand_profiles import CALB_BRAND, GUOXIA_BRAND


@pytest.fixture
def report_sources(saved_cabinet):
    dc, output = saved_cabinet
    data = build_cabinet_report_data(dc, output)
    sld = render_cabinet_sld(output)
    arrangement = render_cabinet_arrangement(output)
    pages = [dict(ac_block_index=s['ac_block_index'], sld_svg=s['svg'].encode(), layout_svg=a['svg'].encode(),
                  sld_png=b'untrusted raster must be ignored', layout_png=None)
             for s, a in zip(sld['pages'], arrangement['pages'])]
    return data, output, pages


def _text(content):
    document = Document(io.BytesIO(content))
    text = '\n'.join(p.text for p in document.paragraphs)
    text += '\n'.join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    return document, text


def test_report_has_saved_values_all_figures_and_no_internal_identifiers(report_sources, monkeypatch):
    from calb_sizing_tool.reporting import cabinet_report as renderer
    original = renderer._add_concept_figure
    figures = []
    def audited(*args, **kwargs):
        figures.append(kwargs)
        return original(*args, **kwargs)
    monkeypatch.setattr(renderer, '_add_concept_figure', audited)
    data, output, pages = report_sources
    before = copy.deepcopy(data)
    content = export_cabinet_report(data, output, pages, brand=CALB_BRAND)
    doc, text = _text(content)
    assert data == before
    assert 'Saved project' in text and '95.00%' in text and '16.16 MWh' in text
    assert '92.70' in text and '90.88' in text and '87.03' in text
    assert '100.00%' in text and '98.00%' in text
    assert len(doc.inline_shapes) >= 5  # one chart + every block's SLD and arrangement
    assert 'private-' not in text and not re.search(r'[A-Za-z]:\\', text)
    assert 'net POI power guarantee' in text and 'NOT FOR CONSTRUCTION' in text
    assert len([section for section in doc.sections if section.page_width > section.page_height]) == 4
    assert len(figures) == 4 and all(f['source_has_status_watermark'] for f in figures)
    # Follow the established engineering report sequence. Cabinet-specific
    # equipment changes the content of a stage, not the stage order.
    assert [paragraph.text for paragraph in doc.paragraphs if paragraph.style.name == 'Heading 1'] == [
        '1 Executive Summary',
        '2 Project Inputs & Assumptions',
        '3 Stage 1 DC Energy Sizing',
        '4 Stage 2 DC Block Configuration',
        '5 Stage 3 Lifetime Degradation & POI Deliverable',
        '6 Stage 4 AC Block Sizing',
        '7 Single Line Diagram',
        '8 Typical AC Block Arrangement (Concept Only)',
        '9 Concept Site Arrangement (Concept Only)',
    ]
    assert [figure['alt_text'].split(' NOT FOR CONSTRUCTION')[0] for figure in figures] == [
        'AC Block 1 electrical connection concept',
        'AC Block 2 electrical connection concept',
        'AC Block 1 diagrammatic cabinet arrangement',
        'AC Block 2 diagrammatic cabinet arrangement',
    ]
    annual = next(table for table in doc.tables if table.cell(0, 0).text == 'Year')
    assessment = next(row for row in annual.rows if row.cells[0].text == '4 *')
    assert [cell.text for cell in assessment.cells] == ['4 *', '92.70', '90.88', '18.23', '16.71', '16.16', '93.00', '87.03']
    from docx.oxml.ns import qn
    assert all(cell._tc.get_or_add_tcPr().find(qn('w:shd')).get(qn('w:fill')) == 'FFF1CC' for cell in assessment.cells)
    assert 'S&C period' in text and 'Operating cycles per year' in text
    assert 'No site envelope, land area, equipment separation or approved route is calculated' in text


def test_white_label_rebuilds_svg_and_contains_no_calb_text(report_sources):
    data, output, pages = report_sources
    data['project_name'] = 'CALB customer project'
    data['dc_products'][0]['block_name'] = 'CALB 418 kWh cabinet'
    content = export_cabinet_report(data, output, pages, brand=GUOXIA_BRAND)
    _, text = _text(content)
    assert 'CALB' not in text.upper()
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        assert all(b'CALB' not in archive.read(name) for name in archive.namelist() if name.endswith('.xml'))


@pytest.mark.parametrize('mutation', ['missing_year', 'duplicate_year', 'wrong_guarantee', 'nan', 'missing_sheet', 'missing_svg'])
def test_incomplete_or_conflicting_report_input_is_rejected(report_sources, mutation):
    data, output, pages = copy.deepcopy(report_sources)
    if mutation == 'missing_year': data['annual_rows'].pop()
    if mutation == 'duplicate_year': data['annual_rows'][-1]['year_index'] = 19
    if mutation == 'wrong_guarantee': data['annual_rows'][4]['is_guarantee_year'] = False
    if mutation == 'nan': data['annual_rows'][2]['poi_usable_energy_mwh'] = float('nan')
    if mutation == 'missing_sheet': pages.pop()
    if mutation == 'missing_svg': pages[0]['sld_svg'] = None
    with pytest.raises(ValueError): export_cabinet_report(data, output, pages)


def test_noncontiguous_membership_prints_actual_members_not_a_false_range(saved_cabinet):
    dc, output = saved_cabinet
    left, right = output['cabinet_group_contract']['groups'][:2]
    left['cabinet_indices'][1], right['cabinet_indices'][0] = right['cabinet_indices'][0], left['cabinet_indices'][1]
    left['cabinet_indices'].sort()
    right['cabinet_indices'].sort()
    data = build_cabinet_report_data(dc, output)
    electrical = render_cabinet_sld(output)['pages']
    spatial = render_cabinet_arrangement(output)['pages']
    pages = [dict(ac_block_index=s['ac_block_index'], sld_svg=s['svg'].encode(), layout_svg=a['svg'].encode())
             for s, a in zip(electrical, spatial)]
    _, text = _text(export_cabinet_report(data, output, pages))
    assert 'C001, C003 to C013 (12 cabinets)' in text
    assert 'C001 to C013 (12 cabinets)' not in text


def test_documented_test_spacing_keeps_source_conditions_and_test_boundary(report_sources):
    data, output, pages = report_sources
    data['installation_basis'] = dict(status='documented', cabinet_dimensions_mm=None,
        clearances=[dict(relation='Illustrative test spacing', value_mm=123,
            evidence_kind='test_configuration', source_title='Fixture report', page=7,
            conditions='Only the described specimen arrangement')], notes=[])
    _, text = _text(export_cabinet_report(data, output, pages))
    assert '123' in text and 'TEST SETUP' in text and 'Fixture report' in text
    assert 'not an installation minimum' in text


def test_unverified_spacing_is_not_promoted_to_a_report_value(report_sources):
    data, output, pages = report_sources
    data['installation_basis'] = dict(status='unverified', cabinet_dimensions_mm={'width': 98765},
        clearances=[dict(value_mm=98765)], notes=[])
    _, text = _text(export_cabinet_report(data, output, pages))
    assert '98765' not in text


def test_documented_dimensions_without_a_source_are_rejected(report_sources):
    data, output, pages = report_sources
    data['installation_basis'] = dict(status='documented', cabinet_dimensions_mm={'width': 1200, 'depth': 1500, 'height': 2300},
        clearances=[], notes=[])
    with pytest.raises(ValueError, match='source'):
        export_cabinet_report(data, output, pages)


def test_stage1_prints_saved_results_without_rederiving_the_dc_requirement(report_sources):
    data, output, pages = copy.deepcopy(report_sources)
    # Deliberate report-boundary values distinguish copying a provided result
    # from recomputing it from the separate POI input or a legacy default.
    data['stage1']['dc_energy_capacity_required_mwh'] = 27.6543
    data['stage1']['dc_power_required_mw'] = 9.8765
    _, text = _text(export_cabinet_report(data, output, pages))
    assert '27.65 MWh' in text and '9.88 MW' in text
    assert 'DC energy capacity required' in text and 'DC equivalent power required' in text


@pytest.mark.parametrize('field,value', [
    ('dc_energy_capacity_required_mwh', None),
    ('dc_energy_capacity_required_mwh', float('nan')),
    ('dc_power_required_mw', float('inf')),
    ('dc_power_required_mw', True),
    ('dc_one_way_efficiency_frac', None),
    ('eff_dc_to_poi_frac', 1.1),
])
def test_missing_or_damaged_stage1_is_rejected_before_document_creation(report_sources, monkeypatch, field, value):
    from calb_sizing_tool.reporting import cabinet_report as renderer
    data, output, pages = copy.deepcopy(report_sources)
    data['stage1'][field] = value
    monkeypatch.setattr(renderer, '_configure_doc', lambda: pytest.fail('Damaged Stage 1 reached DOCX creation'))
    with pytest.raises(ValueError, match='Stage 1'):
        export_cabinet_report(data, output, pages)
