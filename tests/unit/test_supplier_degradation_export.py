from dataclasses import replace
from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook
import pytest

from calb_sizing_tool.services.supplier_degradation_service import parse_supplier_degradation
from calb_sizing_tool.services.supplier_degradation_export_service import export_supplier_review_workbook
from test_supplier_degradation_preview import workbook_bytes


def test_numeric_percentages_missing_points_charts_and_provenance_without_long_ids():
    dataset = parse_supplier_degradation(workbook_bytes())
    output = export_supplier_review_workbook(dataset, metadata={
        'source_sha256': dataset.source_sha256, 'import_id': 'should-not-appear',
        'actor': 'should-not-appear', 'version_label': '=1+1', 'filename': 'supplier.xlsx'})
    book = load_workbook(BytesIO(output))
    lower = book['SOC50_Lower']
    assert lower['B6'].value == 1 and lower['B7'].value == 0.93
    assert lower['B7'].number_format == '0.00%'
    assert lower['B8'].value == 'N/A'
    assert len(lower._charts) == 1
    assert lower._charts[0].y_axis.numFmt.formatCode == '0.00%'
    assert lower._charts[0].series[0].smooth is False
    assert lower.sheet_view.showGridLines is False
    assert lower.freeze_panes == 'B7' and lower.page_setup.fitToWidth == 1
    assert book['Start_here']['A4'].value == 'Version: =1+1'
    assert book['Start_here']['A4'].data_type == 's'
    annual = book['Annual_data']
    assert annual['A5'].value == 0.5 and annual['A5'].number_format == '0.00%'
    assert annual['F6'].value == 0.93 and annual['F6'].number_format == '0.00%'
    assert annual['F7'].value == 'N/A' and annual['H7'].value == 'C7'
    assert annual.freeze_panes == 'A5' and annual.tables
    with ZipFile(BytesIO(output)) as archive:
        text = '\n'.join(archive.read(n).decode() for n in archive.namelist() if n.endswith('.xml'))
        assert dataset.source_sha256 not in text
        assert 'should-not-appear' not in text
        assert '<dispBlanksAs val="gap"' in archive.read('xl/charts/chart1.xml').decode()


def test_extended_cases_are_separate_and_conditional_timing_is_visible():
    dataset = parse_supplier_degradation(workbook_bytes())
    lower = next(c for c in dataset.curves if c.basis == 'lower_limit')
    low_rate = replace(lower, power_rate=0.1, cycles_per_day=2)
    high_frequency = replace(lower, cycles_per_day=3)
    dataset = replace(dataset, curves=(lower, low_rate, high_frequency))
    book = load_workbook(BytesIO(export_supplier_review_workbook(dataset)))
    assert 'SOC50_Lower' in book.sheetnames and 'SOC50_Lower_Extended' in book.sheetnames
    assert len(book['SOC50_Lower']._charts[0].series) == 1
    assert len(book['SOC50_Lower_Extended']._charts) == 2
    rows = list(book['Duty_review'].iter_rows(min_row=5, values_only=True))
    assert [r[7] for r in rows] == ['Fits assumed schedule', 'Exceeds 24 h', 'Exceeds 24 h']
    selected = load_workbook(BytesIO(export_supplier_review_workbook(dataset, [lower])))
    assert 'SOC50_Lower_Extended' not in selected.sheetnames
    assert selected['Annual_data'].max_row == 7
    with pytest.raises(ValueError):
        export_supplier_review_workbook(dataset, [])


def test_full_dataset_all_source_points_are_preserved_and_soc50_comes_first():
    dataset = parse_supplier_degradation(workbook_bytes())
    curves = tuple(replace(c, soc_pct=soc, power_rate=rate, cycles_per_day=freq)
        for soc in [20, 50, 80, 100] for rate in [.1, .25, .33, .5]
        for freq in [.5, 1, 1.5, 2, 2.5, 3] for c in dataset.curves)
    dataset = replace(dataset, curves=curves)
    book = load_workbook(BytesIO(export_supplier_review_workbook(dataset)))
    assert book.sheetnames[1:3] == ['SOC50_Lower', 'SOC50_Mean']
    assert book['Annual_data'].max_row - 4 == sum(len(c.years) for c in curves)
    assert book['Duty_review'].max_row - 4 == len(curves)
    assert len(book['SOC50_Lower']._charts) == 3
    assert len(book['SOC50_Lower_Extended']._charts) == 4
    # Six-frequency low-rate chart has enough printable columns (through Q).
    assert 'Q' in str(book['SOC50_Lower_Extended'].print_area)
