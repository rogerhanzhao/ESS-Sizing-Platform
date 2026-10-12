from io import BytesIO

import pytest
from openpyxl import Workbook
from streamlit.testing.v1 import AppTest

from calb_sizing_tool.services.supplier_degradation_service import parse_supplier_degradation


def workbook_bytes(*, formula=False, lower_exceeds=False):
    book = Workbook()
    book.remove(book.active)
    for name, value in [("Degradation-均值 ", 0.95), ("Degradation-下限值", 0.93)]:
        sheet = book.create_sheet(name)
        sheet['A2'] = 'SOC'
        sheet['C2'] = 0.5
        sheet['C2'].number_format = '0%'
        sheet['A3'] = 'Power Rate'
        sheet['C3'] = 0.25
        sheet['C4'] = 1
        for r, year in enumerate(range(3), 5):
            sheet.cell(r, 2, year)
        sheet['C5'] = 1
        sheet['C6'] = '=0.9' if formula else (0.99 if lower_exceeds and name.endswith('下限值') else value)
        for row in [5, 6]:
            sheet.cell(row, 3).number_format = '0.0%'
    book.create_sheet('RTE')
    book.create_sheet('Calender Degradation')
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def test_source_labels_percentage_units_cells_and_missing_years_are_preserved():
    data = parse_supplier_degradation(workbook_bytes())
    lower = next(c for c in data.curves if c.basis == 'lower_limit')
    assert lower.soc_pct == 50
    assert lower.power_rate == 0.25
    assert lower.cycles_per_day == 1
    assert lower.retention_pct == (100, 93, None)
    assert lower.source_cells == ('C5', 'C6', 'C7')
    assert lower.last_supported_year == 1
    assert data.rte_numeric_values == data.calendar_numeric_values == 0
    assert len(data.source_sha256) == 64


@pytest.mark.parametrize('kwargs, message', [({'formula': True}, 'formulas'), ({'lower_exceeds': True}, 'exceeds')])
def test_unverified_formula_or_inconsistent_lower_limit_is_not_silently_used(kwargs, message):
    with pytest.raises(ValueError, match=message):
        parse_supplier_degradation(workbook_bytes(**kwargs))


def _preview_app():
    from calb_sizing_tool.ui.supplier_degradation_preview import render_supplier_degradation_preview
    render_supplier_degradation_preview()


def test_preview_renders_source_curves_and_preserves_missing_values_without_database_writes(monkeypatch):
    import streamlit as st
    data = workbook_bytes()

    class Upload:
        size = len(data)

        def getvalue(self):
            return data

    monkeypatch.setattr(st, 'file_uploader', lambda *args, **kwargs: Upload())
    app = AppTest.from_function(_preview_app, default_timeout=30).run()
    assert not app.exception
    assert app.selectbox(key='supplier_curve_soc').value == 50
    assert app.selectbox(key='supplier_curve_rate').value == 0.25
    assert any('N/A' in item.value and '93.00%' in item.value for item in app.markdown)
    assert any('Preview only' in item.value for item in app.info)
    assert len(app.get('imgs')) or len(app.get('image'))
    assert any('Download selected reference data' in item.label for item in app.get('download_button'))
