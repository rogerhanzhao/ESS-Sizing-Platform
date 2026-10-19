"""Pure DOCX export for saved independent cabinet group alternatives.

The caller supplies a verified DC record, one persisted AC alternative and all
its verified drawing sheets.  This module never queries a database, resolves a
session or a directory, or recalculates sizing.  The established container
report and drawing engines are deliberately outside this path.
"""

from __future__ import annotations

import datetime as dt
import io
import math
import re
from collections import defaultdict
from typing import Any
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

from calb_sizing_tool.common.render_lock import RENDER_LOCK
from calb_sizing_tool.reporting.brand_profiles import (
    BrandProfile,
    CALB_BRAND,
    assert_brand_clean_docx,
    neutralize_brand_payload,
    neutralize_svg_visible_text,
    require_brand_assets,
)
from calb_sizing_tool.reporting.export_docx import (
    _add_page_number_footer,
    _doc_to_bytes,
    _setup_header,
    set_inline_shape_alt_text,
)
from calb_sizing_tool.services.cabinet_ac_preview_service import validate_cabinet_contract
from calb_sizing_tool.reporting.report_v2 import _add_concept_figure


REPORT_VERSION = "Cabinet technical proposal V2"
REPORT_CHAPTERS = (
    "1 Executive Summary",
    "2 Project Inputs & Assumptions",
    "3 Stage 1 DC Energy Sizing",
    "4 Stage 2 DC Block Configuration",
    "5 Stage 3 Lifetime Degradation & POI Deliverable",
    "6 Stage 4 AC Block Sizing",
    "7 Single Line Diagram",
    "8 Typical AC Block Arrangement (Concept Only)",
    "9 Concept Site Arrangement (Concept Only)",
)
_UUID = re.compile(r"\b[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}\b", re.I)
_LOCAL_PATH = re.compile(r"(?:[A-Za-z]:[\\/]|file://|\\\\)", re.I)
_ANNUAL_NUMBERS = (
    "soh_display_pct", "soh_absolute_pct", "dc_gross_capacity_mwh",
    "dc_usable_mwh", "dc_rte_pct", "system_rte_pct", "poi_usable_energy_mwh",
)


def _integer(value: Any, label: str, *, minimum: int = 0) -> int:
    if type(value) is not int or value < minimum:
        raise ValueError(f"{label} must be an integer of at least {minimum}")
    return value


def _number(value: Any, label: str, *, minimum: float = 0) -> float:
    if isinstance(value, bool):
        raise ValueError(f"{label} must be a finite number")
    try:
        number = float(value)
    except (ValueError, TypeError, OverflowError) as exc:
        raise ValueError(f"{label} must be a finite number") from exc
    if not math.isfinite(number) or number < minimum:
        raise ValueError(f"{label} must be a finite number of at least {minimum}")
    return number


def _same(value: Any, expected: float, label: str, *, tolerance: float = .00051) -> None:
    if not math.isclose(_number(value, label), expected, rel_tol=0, abs_tol=tolerance):
        raise ValueError(f"{label} conflicts with the saved cabinet alternative")


def _display(value: Any, default: str = "To be confirmed") -> str:
    """Only customer-facing names are printed; internal identifiers stay private."""
    text = " ".join(str(value or "").split())
    if _LOCAL_PATH.search(text):
        return default
    text = _UUID.sub("", text).strip(" /|_- ")
    return text or default


def _fmt(value: Any, suffix: str = "", decimals: int = 2) -> str:
    return "To be confirmed" if value is None else f"{float(value):,.{decimals}f}{suffix}"


def _cabinet_membership(indices: list[int]) -> str:
    """Compact only real consecutive runs; never imply an undrawn cabinet."""
    members = sorted(indices)
    runs: list[tuple[int, int]] = []
    start = previous = members[0]
    for member in members[1:]:
        if member != previous + 1:
            runs.append((start, previous))
            start = member
        previous = member
    runs.append((start, previous))
    text = ", ".join(f"C{start:03d}" if start == end else f"C{start:03d} to C{end:03d}" for start, end in runs)
    return f"{text} ({len(members)} cabinets)"


def _validate(report_data: dict, ac_output: dict, diagram_pages: list[dict]) -> tuple[list[dict], dict[int, list[dict]]]:
    if not isinstance(report_data, dict) or not isinstance(ac_output, dict):
        raise ValueError("Saved report data and AC output are required")
    stage1 = report_data.get("stage1")
    if not isinstance(stage1, dict):
        raise ValueError("The saved Stage 1 DC sizing result is required")
    for field in ("dc_energy_capacity_required_mwh", "dc_power_required_mw"):
        _number(stage1.get(field), f"Saved Stage 1 {field}", minimum=.000001)
    for field in ("eff_dc_cables_frac", "eff_pcs_frac", "eff_mvt_frac", "eff_ac_cables_sw_rmu_frac",
                  "eff_hvt_others_frac", "eff_dc_to_poi_frac", "dod_frac",
                  "dc_round_trip_efficiency_frac", "dc_one_way_efficiency_frac"):
        if _number(stage1.get(field), f"Saved Stage 1 {field}", minimum=.000001) > 1:
            raise ValueError(f"Saved Stage 1 {field} must be a fraction no greater than one")
    if _number(stage1.get("sc_loss_pct"), "Saved Stage 1 S&C loss") > 100:
        raise ValueError("Saved Stage 1 S&C loss must be a percentage no greater than 100")
    if ac_output.get("dc_unit_kind") != "independent_cabinet_group":
        raise ValueError("A saved independent cabinet group alternative is required")
    plan = ac_output.get("cabinet_group_contract")
    validate_cabinet_contract(plan, allocated=True)
    if ac_output.get("transformer_topology") not in {"two_winding", "three_winding"} or ac_output.get("transformer_topology_confirmation") != "confirmed_by_user":
        raise ValueError("The saved transformer LV topology must be explicitly confirmed")
    count = _integer(report_data.get("physical_cabinet_count"), "Physical cabinet count", minimum=1)
    if count != plan["physical_cabinet_count"] or count != ac_output.get("physical_cabinet_count"):
        raise ValueError("Physical cabinet count conflicts with the saved cabinet alternative")
    _same(report_data.get("dc_total_mwh"), plan["total_nameplate_mwh"], "Installed DC nameplate")
    _same(ac_output.get("dc_total_mwh"), plan["total_nameplate_mwh"], "AC source DC nameplate")

    number_of_blocks = _integer(ac_output.get("num_blocks"), "AC Block count", minimum=1)
    ac_rows = ac_output.get("ac_block_rows")
    if not isinstance(ac_rows, list) or len(ac_rows) != number_of_blocks:
        raise ValueError("Complete actual AC Block equipment rows are required")
    by_block: dict[int, list[dict]] = defaultdict(list)
    for group in plan["groups"]:
        if group["ac_block_index"] > number_of_blocks:
            raise ValueError("Cabinet group is outside the saved AC Block schedule")
        by_block[group["ac_block_index"]].append(group)
    pcs_total, pcs_power = 0, 0.0
    for index, row in enumerate(ac_rows, 1):
        pcs_count = _integer(row.get("pcs_count"), "PCS count", minimum=1)
        pcs_kw = _number(row.get("pcs_kw"), "PCS unit rating", minimum=.000001)
        groups = by_block[index]
        if pcs_count != len(groups) or row.get("dc_blocks") != len(groups):
            raise ValueError("Exactly one intact independent cabinet group per PCS is required")
        if sorted(g["pcs_index"] for g in groups) != list(range(1, pcs_count + 1)):
            raise ValueError("Saved group to PCS assignment is incomplete")
        pcs_total += pcs_count
        pcs_power += pcs_count * pcs_kw / 1000
    if ac_output.get("pcs_count_by_block") != [row["pcs_count"] for row in ac_rows]:
        raise ValueError("PCS equipment mirrors conflict with the saved AC schedule")
    if ac_output.get("pcs_count_total", pcs_total) != pcs_total:
        raise ValueError("Total PCS count conflicts with the saved AC schedule")
    _same(ac_output.get("total_ac_mw", pcs_power), pcs_power, "PCS nameplate total", tolerance=.000001)

    products = report_data.get("dc_products")
    if not isinstance(products, list) or not products:
        raise ValueError("A saved DC product equipment schedule is required")
    product_count, product_energy = 0, 0.0
    for product in products:
        product_count += _integer(product.get("count"), "DC product count", minimum=1)
        _number(product.get("unit_capacity_mwh"), "DC product unit capacity", minimum=.000001)
        subtotal = _number(product.get("subtotal_mwh"), "DC product subtotal", minimum=.000001)
        _same(subtotal, product["count"] * float(product["unit_capacity_mwh"]), "DC product subtotal")
        product_energy += subtotal
    if product_count != count:
        raise ValueError("DC product schedule does not cover all physical cabinets")
    _same(product_energy, plan["total_nameplate_mwh"], "DC product equipment total")

    life = _integer(report_data.get("project_life_years"), "Project life", minimum=1)
    guarantee_year = _integer(report_data.get("poi_guarantee_year"), "Energy assessment year")
    if guarantee_year > life:
        raise ValueError("The energy assessment year is outside the saved project life")
    _number(report_data.get("poi_power_mw"), "POI power requirement", minimum=.000001)
    _number(report_data.get("poi_energy_mwh"), "POI energy requirement", minimum=.000001)
    annual = report_data.get("annual_rows")
    if not isinstance(annual, list) or len(annual) != life + 1:
        raise ValueError("A complete saved annual performance schedule is required")
    years = [_integer(row.get("year_index"), "Annual year") for row in annual]
    if sorted(years) != list(range(life + 1)):
        raise ValueError("Annual performance years must cover COD through the project life exactly once")
    annual = sorted(annual, key=lambda row: row["year_index"])
    for row in annual:
        for field in _ANNUAL_NUMBERS:
            number = _number(row.get(field), f"Annual {field}")
            if field.endswith("_pct") and number > 100.000001:
                raise ValueError(f"Annual {field} must be a percentage from zero to 100")
        for flag in ("is_guarantee_year", "meets_poi_req"):
            if type(row.get(flag)) is not bool:
                raise ValueError(f"Annual {flag} must be an explicit saved boolean")
        if row["is_guarantee_year"] != (row["year_index"] == guarantee_year):
            raise ValueError("Saved annual assessment flag conflicts with the selected assessment year")
        if "dc_nameplate_bol_mwh" in row:
            _same(row["dc_nameplate_bol_mwh"], plan["total_nameplate_mwh"], "Annual installed DC nameplate")
    for field in ("dod_pct", "dc_round_trip_efficiency_pct"):
        if report_data.get(field) is not None and _number(report_data[field], field) > 100:
            raise ValueError(f"{field} must be a percentage from zero to 100")

    if not isinstance(diagram_pages, list) or len(diagram_pages) != number_of_blocks:
        raise ValueError("Both verified drawing sheets are required for every actual AC Block")
    if sorted(_integer(p.get("ac_block_index"), "Drawing AC Block", minimum=1) for p in diagram_pages) != list(range(1, number_of_blocks + 1)):
        raise ValueError("Drawing sheets must cover every actual AC Block exactly once")
    for page in diagram_pages:
        for kind in ("sld", "layout"):
            svg = page.get(f"{kind}_svg")
            if not isinstance(svg, bytes) or not svg:
                raise ValueError(f"The {kind} SVG source is required for brand safe export")
            try:
                root = ET.fromstring(svg)
            except ET.ParseError as exc:
                raise ValueError(f"The {kind} SVG is not readable") from exc
            if root.tag.rsplit("}", 1)[-1] != "svg":
                raise ValueError(f"The {kind} source is not an SVG drawing")
            visible = " ".join(root.itertext()).upper()
            if "NOT FOR CONSTRUCTION" not in visible or "CONCEPT" not in visible:
                raise ValueError(f"The {kind} SVG must retain the visible concept status warning")
    return annual, dict(by_block)


def _configure_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(.8)
    section.top_margin = Inches(.8)
    section.bottom_margin = Inches(.7)
    section.header_distance = Inches(.25)
    section.footer_distance = Inches(.25)
    for name in ("Normal", "Body Text", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(11 if name in {"Normal", "Body Text"} else 10)
        style.paragraph_format.space_after = Pt(7)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.08
    doc.styles["Title"].font.size = Pt(28)
    doc.styles["Title"].font.underline = False
    title_properties = doc.styles["Title"].element.get_or_add_pPr()
    for border in list(title_properties.findall(qn("w:pBdr"))):
        title_properties.remove(border)
    doc.styles["Heading 1"].font.size = Pt(17)
    doc.styles["Heading 2"].font.size = Pt(13)
    doc.styles["Subtitle"].font.size = Pt(14)
    doc.styles["Caption"].font.size = Pt(9)
    doc.styles["Caption"].font.bold = False
    doc.styles["Caption"].font.italic = False
    return doc


def _table(doc, headers, rows, *, widths=None, font_size=10, numeric_columns=(), highlight_rows=()):
    """Match the established report palette with explicit readable geometry."""
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    available = (doc.sections[-1].page_width - doc.sections[-1].left_margin - doc.sections[-1].right_margin) / 914400
    fractions = widths or [1 / len(headers)] * len(headers)
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = str(header)
    for values in rows:
        for cell, value in zip(table.add_row().cells, values):
            cell.text = str(value)
    for i, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))
        for j, cell in enumerate(row.cells):
            cell.width = Inches(available * fractions[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for edge, amount in (("top", 65), ("bottom", 65), ("left", 80), ("right", 80)):
                element = OxmlElement(f"w:{edge}")
                element.set(qn("w:w"), str(amount))
                element.set(qn("w:type"), "dxa")
                margins.append(element)
            tc_pr.append(margins)
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "bottom", "left", "right", "insideH", "insideV"):
                element = OxmlElement(f"w:{edge}")
                for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    element.set(qn(f"w:{key}"), value)
                borders.append(element)
            tc_pr.append(borders)
            shading = OxmlElement("w:shd")
            highlighted = i > 0 and i - 1 in highlight_rows
            shading.set(qn("w:fill"), "1F4E79" if i == 0 else "FFF1CC" if highlighted else "F2F6FA" if i % 2 == 0 else "FFFFFF")
            tc_pr.append(shading)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                if j in numeric_columns:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)
                    if i == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                    elif highlighted:
                        run.font.bold = True
    for j, fraction in enumerate(fractions):
        table.columns[j].width = Inches(available * fraction)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def _page(doc, title: str):
    doc.add_page_break()
    _heading(doc, title)


def _bookmark_name(title: str) -> str:
    return "cab_" + re.sub(r"[^a-z0-9]", "_", title.lower())[:35]


def _bookmark(paragraph, title: str):
    name = _bookmark_name(title)
    # Stable unique identifiers for this report's known section titles.
    identifier = str(sum((i + 1) * ord(char) for i, char in enumerate(name)))
    start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
    start.set(qn("w:id"), identifier)
    start.set(qn("w:name"), name)
    end.set(qn("w:id"), identifier)
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _heading(doc, title: str):
    paragraph = doc.add_heading(title, level=1)
    _bookmark(paragraph, title)
    return paragraph


def _contents_link(doc, title: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), _bookmark_name(title))
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = title
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def _small_section_break(doc):
    """Keep the section marker from spilling a full-page drawing to a blank page."""
    marker = doc.paragraphs[-1]
    marker.paragraph_format.space_before = Pt(0)
    marker.paragraph_format.space_after = Pt(0)
    marker.paragraph_format.line_spacing = Pt(1)
    marker.add_run().font.size = Pt(1)


def _add_lifetime_chart(doc, annual, target, guarantee_year):
    """Embed only the chart of saved annual energy/SOH data, never a drawing."""
    content = _annual_chart(annual, target, guarantee_year)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    shape = paragraph.add_run().add_picture(io.BytesIO(content), width=Inches(6.7))
    set_inline_shape_alt_text(shape, 'Saved annual POI usable energy and relative versus absolute SOH from COD')
    return shape


def _annual_chart(annual: list[dict], target: float, guarantee_year: int) -> bytes:
    """Plot saved records verbatim; no degradation or sizing calculation."""
    from matplotlib.figure import Figure

    with RENDER_LOCK:
        figure = Figure(figsize=(7.4, 5.2), dpi=180)
        energy, soh = figure.subplots(2, 1, sharex=True)
        years = [row["year_index"] for row in annual]
        energy.plot(years, [row["poi_usable_energy_mwh"] for row in annual], color="#1f4e79", linewidth=2.3, marker="o", markersize=3, label="Saved model POI usable energy")
        energy.axhline(target, color="#b24e26", linestyle="--", linewidth=1.4, label="POI energy requirement")
        energy.axvline(guarantee_year, color="#566673", linestyle=":", linewidth=1, label=f"Selected assessment year {guarantee_year}")
        energy.set_ylabel("POI usable energy (MWh)")
        energy.legend(fontsize=8, loc="best", frameon=False)
        soh.plot(years, [row["soh_display_pct"] for row in annual], color="#1f4e79", linewidth=2, label="SOH relative to COD")
        soh.plot(years, [row["soh_absolute_pct"] for row in annual], color="#b24e26", linewidth=1.7, linestyle="--", label="Absolute SOH")
        soh.set_ylabel("SOH (%)")
        soh.set_xlabel("Year from commercial operation date")
        soh.legend(fontsize=8, loc="best", frameon=False)
        soh.set_xticks(years if len(years) <= 25 else years[::max(1, len(years) // 20)])
        for axis in (energy, soh):
            axis.grid(axis="y", color="#dce3ea", linewidth=.6)
            axis.spines[["top", "right"]].set_visible(False)
            axis.tick_params(labelsize=8)
        figure.tight_layout(pad=1.3)
        buffer = io.BytesIO()
        figure.savefig(buffer, format="png", facecolor="white", dpi=180)
        return buffer.getvalue()


def _rasterize_svg(svg: bytes, brand: BrandProfile) -> bytes:
    """Always regenerate from audited SVG, including when a PNG was supplied."""
    import cairosvg

    safe = neutralize_svg_visible_text(svg, brand)
    root = ET.fromstring(safe)
    # Raster images can carry an unauditable publisher mark; new vector-only
    # cabinet renderers have no such elements or external resources.
    for element in root.iter():
        tag = element.tag.rsplit("}", 1)[-1].lower()
        if tag in {"image", "script", "foreignobject"}:
            raise ValueError("Cabinet report drawings must contain auditable vector content only")
        for key, value in element.attrib.items():
            if key.rsplit("}", 1)[-1].lower() == "href" and not value.startswith("#"):
                raise ValueError("External drawing resources are not permitted in the cabinet report")
    with RENDER_LOCK:
        return cairosvg.svg2png(bytestring=safe, output_width=3000, background_color="white")


def _diagram_section(doc, svg: bytes, brand: BrandProfile, description: str, *, chapter_title: str | None = None):
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    _small_section_break(doc)
    # Preserve the original A3 engineering drawing scale instead of shrinking
    # cabinet identifiers below readable print size on a Letter page.
    section.page_width, section.page_height = Mm(420), Mm(297)
    section.left_margin = section.right_margin = Inches(.5)
    section.top_margin = Inches(.7)
    section.bottom_margin = Inches(.4)
    section.header_distance = Inches(.15)
    section.footer_distance = Inches(.15)
    if chapter_title:
        paragraph = _heading(doc, chapter_title)
        # The drawing pages reserve a compact heading line above a full A3
        # figure. Keep Heading 1 semantics without inheriting its larger body
        # chapter spacing, which can push the first drawing to another page.
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            run.font.size = Pt(13)
    else:
        paragraph = doc.add_heading(description.split(" NOT FOR CONSTRUCTION")[0], level=2)
        paragraph.paragraph_format.space_after = Pt(4)
    png = _rasterize_svg(svg, brand)
    from PIL import Image
    with Image.open(io.BytesIO(png)) as image:
        width, height = image.size
    usable_width = (section.page_width - section.left_margin - section.right_margin) / 914400
    usable_height = (section.page_height - section.top_margin - section.bottom_margin) / 914400 - .7
    _add_concept_figure(doc, png, width=Inches(min(usable_width, usable_height * width / height)),
                        source_has_status_watermark=True, alt_text=description)
    paragraph = doc.paragraphs[-1]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(0)
    _bookmark(doc.paragraphs[-1], description.split(" NOT FOR CONSTRUCTION")[0])


def _portrait_section(doc, title: str):
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    _small_section_break(doc)
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(.8)
    section.top_margin = Inches(.8)
    section.bottom_margin = Inches(.7)
    section.header_distance = section.footer_distance = Inches(.25)
    _heading(doc, title)


def _installation_basis(doc, basis: Any):
    """Present a source-bound installation basis without promoting test geometry."""
    doc.add_heading("Installation basis and conditions", level=2)
    if not isinstance(basis, dict) or basis.get("status") != "documented":
        doc.add_paragraph(
            "Manufacturer installation requirements and source test configurations have not been verified for this export. "
            "Cabinet dimensions, installation clearances and separation distances remain to be confirmed. "
            "No zero clearance or minimum separation is implied by the concept arrangement."
        )
        return
    doc.add_paragraph(
        "The following values are a documented reference basis. Their stated conditions and evidence type must remain attached to each value. "
        "Test configuration distances describe the tested setup and do not establish minimum installation clearances or general site approval."
    )
    dimensions = basis.get("cabinet_dimensions_mm")
    if dimensions is not None:
        if not isinstance(dimensions, dict) or set(dimensions) != {"width", "depth", "height"}:
            raise ValueError("Documented cabinet dimensions require width, depth and height in mm")
        dimension_source = basis.get("dimension_source")
        if (not isinstance(dimension_source, dict)
                or not isinstance(dimension_source.get("source_title"), str)
                or not dimension_source["source_title"].strip()
                or type(dimension_source.get("page")) is not int
                or dimension_source["page"] < 1
                or ("revision" in dimension_source and not isinstance(dimension_source["revision"], str))):
            raise ValueError("Documented cabinet dimensions require an explicit source title and positive page")
        values = []
        for field, label in (("width", "Width"), ("depth", "Depth"), ("height", "Height")):
            values.append(f"{label} {_fmt(_number(dimensions[field], label, minimum=.000001), ' mm', 0)}")
        doc.add_paragraph("Manufacturer cabinet dimensions: " + "; ".join(values) + ".")
        doc.add_paragraph(
            f"Dimension source: {_display(dimension_source['source_title'])}, page {dimension_source['page']}"
            + (f", revision {_display(dimension_source['revision'])}" if dimension_source.get("revision") else "") + ".", style="Caption"
        )
    clearances = basis.get("clearances") or []
    if not isinstance(clearances, list):
        raise ValueError("Installation clearances must be a source-bound list")
    if dimensions is None and not clearances:
        raise ValueError("A documented installation basis requires dimensional or clearance evidence")
    for item in clearances:
        if not isinstance(item, dict) or item.get("evidence_kind") not in {"manufacturer_requirement", "test_configuration"}:
            raise ValueError("Each installation distance requires an explicit manufacturer or test evidence kind")
        value = _number(item.get("value_mm"), "Documented installation distance")
        if (not isinstance(item.get("relation"), str) or not item["relation"].strip()
                or not isinstance(item.get("source_title"), str) or not item["source_title"].strip()
                or type(item.get("page")) is not int or item["page"] < 1):
            raise ValueError("Each installation distance requires its relation, source title and page")
        kind = "MANUFACTURER REQUIREMENT" if item["evidence_kind"] == "manufacturer_requirement" else "TEST SETUP ONLY"
        paragraph = doc.add_paragraph()
        paragraph.add_run(f"{kind}: {_display(item['relation'])} — {_fmt(value, ' mm', 0)}. ").bold = True
        if item.get("conditions"):
            conditions = item["conditions"]
            paragraph.add_run(_display("; ".join(map(str, conditions)) if isinstance(conditions, list) else conditions) + " ")
        if item["evidence_kind"] == "test_configuration":
            paragraph.add_run("This is a tested configuration, not an installation minimum. ")
        doc.add_paragraph(f"Source: {_display(item['source_title'])}, page {_display(item['page'])}.", style="Caption")
    notes = basis.get("notes") or []
    if isinstance(notes, str):
        notes = [notes]
    for note in notes:
        doc.add_paragraph(_display(note))


def _brand_headers(doc, brand: BrandProfile):
    _setup_header(doc, brand.header_title, logo_path=brand.logo_path,
                  header_lines=list(brand.header_lines), footer_lines=list(brand.footer_lines))
    _add_page_number_footer(doc)
    # The shared header uses the portrait content width. Extend only this new
    # report's generated header tables to each drawing sheet's width.
    for section in doc.sections:
        available = section.page_width - section.left_margin - section.right_margin
        for table in section.header.tables:
            table.autofit = False
            for row in table.rows:
                for index, cell in enumerate(row.cells):
                    cell.width = int(available * (.4 if index == 0 else .6))
            for index, column in enumerate(table.columns):
                column.width = int(available * (.4 if index == 0 else .6))


def export_cabinet_report(report_data: dict, ac_output: dict, diagram_pages: list[dict], *, brand: BrandProfile | None = None) -> bytes:
    """Export the complete selected cabinet alternative as an English DOCX.

    All per-year figures originate in ``annual_rows``. Source UUIDs, filesystem
    paths and raw JSON are never printed. A formal capability or construction
    status cannot be selected by this concept report generator.
    """
    annual, by_block = _validate(report_data, ac_output, diagram_pages)
    brand = brand or CALB_BRAND
    require_brand_assets(brand)
    data = neutralize_brand_payload(report_data, brand)
    output = neutralize_brand_payload(ac_output, brand)
    plan = output["cabinet_group_contract"]
    doc = _configure_doc()
    doc.core_properties.title = "Cabinet Battery Energy Storage Technical Proposal"
    doc.core_properties.subject = "Saved sizing results and concept electrical arrangement"
    doc.core_properties.author = brand.company_legal_name
    doc.core_properties.keywords = "concept; not for construction; cabinet BESS"
    doc.core_properties.comments = REPORT_VERSION
    count, total = data["physical_cabinet_count"], data["dc_total_mwh"]
    guarantee_year = data["poi_guarantee_year"]
    assessed = annual[guarantee_year]
    rows = output["ac_block_rows"]
    pcs_total = sum(row["pcs_count"] for row in rows)
    pcs_mw = sum(row["pcs_count"] * float(row["pcs_kw"]) / 1000 for row in rows)

    doc.add_paragraph("Battery Energy Storage System Technical Proposal", style="Title")
    doc.add_paragraph(_display(data.get("project_name"), "Battery Energy Storage Project"), style="Subtitle")
    doc.add_paragraph("CONCEPT / NOT FOR CONSTRUCTION").runs[0].bold = True
    identity = [
        ["Project", _display(data.get("project_name"))],
        ["Case", _display(data.get("case_name"))],
        ["Configuration", f"{count} cabinets / {len(plan['groups'])} independent DC groups / {len(rows)} AC Blocks"],
        ["POI requirement", f"{_fmt(data['poi_power_mw'], ' MW')} / {_fmt(data['poi_energy_mwh'], ' MWh')}"],
        ["Installed DC nameplate at BOL", _fmt(total, " MWh", 3)],
        ["Energy assessment point", f"Year {guarantee_year} / {_fmt(assessed['poi_usable_energy_mwh'], ' MWh')} modelled POI energy"],
    ]
    if data.get("ac_alternative_label"):
        identity.append(["AC alternative", _display(data["ac_alternative_label"])])
    identity.extend([["Issue date", dt.date.today().isoformat()], ["Document status", "Concept technical proposal"], ["Prepared by", brand.company_legal_name]])
    _table(doc, ["Document information", "Selected basis"], identity, widths=[.34, .66])
    doc.add_paragraph("This proposal follows the DC energy sizing, DC configuration, lifetime performance and AC Block engineering sequence. All results and equipment assignments refer to the selected saved alternative.")
    doc.add_paragraph(brand.confidentiality_notice, style="Caption")
    doc.add_heading("Contents", level=2)
    for title in REPORT_CHAPTERS:
        _contents_link(doc, title)

    _page(doc, REPORT_CHAPTERS[0])
    doc.add_paragraph(
        f"The selected configuration provides {_fmt(assessed['poi_usable_energy_mwh'], ' MWh')} modelled usable energy at the POI in Year {guarantee_year} "
        f"against {_fmt(data['poi_energy_mwh'], ' MWh')} required. The saved energy assessment is "
        f"{'requirement met' if assessed['meets_poi_req'] else 'requirement not met'}."
    )
    _table(doc, ["Metric", "Selected configuration"], [
        ["POI power requirement", _fmt(data["poi_power_mw"], " MW")],
        ["POI energy requirement", _fmt(data["poi_energy_mwh"], " MWh")],
        ["Energy assessment year from COD", guarantee_year],
        ["POI usable energy at assessment year", _fmt(assessed["poi_usable_energy_mwh"], " MWh")],
        ["Saved energy assessment", "Requirement met" if assessed["meets_poi_req"] else "Requirement not met"],
        ["Physical battery cabinets", count],
        ["DC nameplate at BOL", _fmt(total, " MWh", 3)],
        ["Independent DC groups", len(plan["groups"])],
        ["AC Blocks / PCS units", f"{len(rows)} / {pcs_total}"],
        ["Installed PCS nameplate", _fmt(pcs_mw, " MW")],
    ], widths=[.59, .41])
    doc.add_paragraph("One intact independent DC group feeds one PCS. The PCS nameplate is distinct from net POI capability; this report does not establish a net POI power guarantee or a contractual performance warranty. The engineering figures are concept drawings.")

    _page(doc, REPORT_CHAPTERS[1])
    _table(doc, ["Parameter", "Saved project input"], [
        ["POI power requirement", _fmt(data["poi_power_mw"], " MW")],
        ["POI usable energy requirement", _fmt(data["poi_energy_mwh"], " MWh")],
        ["Selected energy assessment year", f"Year {guarantee_year}"],
        ["Project life", f"{data['project_life_years']} years"],
        ["Operating cycles per year", _fmt(data.get("cycles_per_year"), " cycles/year", 0)],
        ["S&C period", _fmt(data.get("sc_time_months"), " months", 0)],
        ["Depth of discharge", _fmt(data.get("dod_pct"), "%")],
        ["DC round trip efficiency input", _fmt(data.get("dc_round_trip_efficiency_pct"), "%")],
        ["POI nominal voltage", _fmt(data.get("poi_nominal_voltage_kv"), " kV")],
        ["Grid frequency", _fmt(data.get("poi_frequency_hz"), " Hz")],
    ], widths=[.56, .44])
    doc.add_paragraph("The following stages reproduce the selected saved results. Energy conversion efficiencies are the saved one way DC to POI values; no new auxiliary load, dispatch, augmentation or installation assumption is introduced by this report.")

    _page(doc, REPORT_CHAPTERS[2])
    stage1 = data["stage1"]
    _table(doc, ["Saved Stage 1 result", "Value"], [
        ["DC energy capacity required", _fmt(stage1["dc_energy_capacity_required_mwh"], " MWh")],
        ["DC equivalent power required", _fmt(stage1["dc_power_required_mw"], " MW")],
    ], widths=[.64, .36])
    doc.add_paragraph("These DC requirements are copied directly from the saved Stage 1 sizing result. Report preparation does not recalculate the requirement, alter its efficiency chain or select another degradation curve.")
    _table(doc, ["Stage 1 input boundary", "Saved value"], [
        ["S&C loss", _fmt(stage1["sc_loss_pct"], "%")],
        ["Depth of discharge", _fmt(stage1["dod_frac"] * 100, "%")],
        ["DC RTE round trip", _fmt(stage1["dc_round_trip_efficiency_frac"] * 100, "%")],
        ["DC one way discharge efficiency", _fmt(stage1["dc_one_way_efficiency_frac"] * 100, "%")],
    ], widths=[.64, .36])
    doc.add_heading("3.1 Efficiency Chain One Way DC to POI", level=2)
    _table(doc, ["Component", "Saved efficiency"], [
        ["Total one way efficiency", _fmt(stage1["eff_dc_to_poi_frac"] * 100, "%")],
        ["DC cables", _fmt(stage1["eff_dc_cables_frac"] * 100, "%")],
        ["PCS", _fmt(stage1["eff_pcs_frac"] * 100, "%")],
        ["MV transformer", _fmt(stage1["eff_mvt_frac"] * 100, "%")],
        ["RMU switchgear and AC cables", _fmt(stage1["eff_ac_cables_sw_rmu_frac"] * 100, "%")],
        ["HVT and others", _fmt(stage1["eff_hvt_others_frac"] * 100, "%")],
    ], widths=[.64, .36])
    doc.add_paragraph("The saved energy model boundary is distinct from final equipment acceptance and the project auxiliary load balance. These remain subject to the agreed equipment data and performance test boundary.", style="Caption")

    _page(doc, REPORT_CHAPTERS[3])
    doc.add_paragraph(f"The selected cabinet configuration contains {count} cabinets with {_fmt(total, ' MWh', 3)} DC nameplate at BOL. Physical cabinet quantity and independent electrical group quantity are shown separately.")
    _table(doc, ["Battery equipment", "Quantity", "Unit MWh", "Total MWh"], [
        [_display(product.get("block_name") or product.get("block_code"), "418 kWh battery cabinet"), product["count"], _fmt(product["unit_capacity_mwh"], decimals=3), _fmt(product["subtotal_mwh"], decimals=3)]
        for product in data["dc_products"]
    ], widths=[.52, .12, .17, .19], numeric_columns=(1, 2, 3))
    doc.add_heading("4.1 Independent Group Connection Schedule", level=2)
    _table(doc, ["Group", "Physical cabinet membership", "DC MWh", "AC Block", "PCS"], [
        [f"G{g['group_index']:02d}", _cabinet_membership(g["cabinet_indices"]),
         _fmt(g["nameplate_mwh"], decimals=3), f"AC {g['ac_block_index']:02d}", f"{g['ac_block_index']:02d}-{g['pcs_index']:02d}"]
        for g in sorted(plan["groups"], key=lambda g: g["group_index"])
    ], widths=[.12, .42, .16, .16, .14], numeric_columns=(2,))
    doc.add_paragraph("Each independent group contains 1 to 12 physical cabinets and connects intact to one PCS. The schedule lists the actual saved cabinet membership and PCS assignment; no DC tie between independent groups is assumed.")

    _page(doc, REPORT_CHAPTERS[4])
    curve_basis = data.get("stage3_meta") or {}
    _table(doc, ["Lifetime model parameter", "Saved value"], [
        ["Effective DC C rate", _fmt(curve_basis.get("effective_c_rate"), " C", 3)],
        ["Selected SOH reference curve C rate", _fmt(curve_basis.get("chosen_soh_c_rate"), " C", 3)],
        ["Selected SOH reference curve cycles per year", _fmt(curve_basis.get("chosen_soh_cycles_per_year"), " cycles/year", 0)],
        ["Selected RTE reference curve C rate", _fmt(curve_basis.get("chosen_rte_c_rate"), " C", 3)],
    ], widths=[.66, .34])
    doc.add_paragraph("Annual DC RTE and system RTE are the saved Stage 3 outputs. They can differ from the Stage 1 reference input; the report preserves both modelling bases.", style="Caption")
    _add_lifetime_chart(doc, annual, float(data["poi_energy_mwh"]), guarantee_year)
    doc.add_paragraph(
        f"At Year {guarantee_year}, the saved model reports {_fmt(assessed['poi_usable_energy_mwh'], ' MWh')} usable POI energy "
        f"against {_fmt(data['poi_energy_mwh'], ' MWh')} required. Saved energy assessment: "
        f"{'requirement met' if assessed['meets_poi_req'] else 'requirement not met'}. The saved curve selection and annual results are unchanged."
    )
    doc.add_paragraph(
        f"SOH relative to COD is {_fmt(annual[0]['soh_display_pct'], '%')} at Year 0; absolute SOH is {_fmt(annual[0]['soh_absolute_pct'], '%')} at the same checkpoint. "
        "The two curves use different reference bases. Modelled energy is distinct from final equipment operating limits and contractual warranty terms.", style="Caption"
    )
    doc.add_page_break()
    doc.add_heading("5.1 Annual Energy and SOH Schedule", level=2)
    doc.add_paragraph(f"The highlighted row marked * is the selected Year {guarantee_year} energy assessment. Other years are reference projections, not additional contractual pass or fail checkpoints.")
    _table(doc, ["Year", "SOH vs\nCOD (%)", "Absolute\nSOH (%)", "DC gross\nMWh", "DC usable\nMWh", "POI usable\nMWh", "DC RTE\n(%)", "System RTE\n(%)"], [
        [f"{row['year_index']}{' *' if row['is_guarantee_year'] else ''}", _fmt(row["soh_display_pct"]), _fmt(row["soh_absolute_pct"]), _fmt(row["dc_gross_capacity_mwh"]), _fmt(row["dc_usable_mwh"]), _fmt(row["poi_usable_energy_mwh"]), _fmt(row["dc_rte_pct"]), _fmt(row["system_rte_pct"])]
        for row in annual
    ], widths=[.07, .12, .12, .14, .14, .15, .12, .14], font_size=9, numeric_columns=tuple(range(1, 8)), highlight_rows=(guarantee_year,))
    doc.add_paragraph("SOH vs COD uses the commercial operation date baseline. Absolute SOH retains the saved absolute capacity reference. DC gross and DC usable energy are at the DC plane; POI usable energy includes the saved conversion boundary. Percentage values are displayed to two decimal places without changing the stored values or assessment.", style="Caption")

    _page(doc, REPORT_CHAPTERS[5])
    doc.add_paragraph(
        f"The selected alternative uses {pcs_total} PCS with {_fmt(pcs_mw, ' MW')} combined nameplate capacity. "
        f"The POI power requirement is {_fmt(data['poi_power_mw'], ' MW')}. Equality of these two numbers would not establish a net POI power guarantee."
    )
    _table(doc, ["AC Block", "Cabinets", "DC groups", "DC MWh", "PCS configuration", "PCS MW"], [
        [f"AC {index:02d}", sum(g["cabinet_count"] for g in by_block[index]), len(by_block[index]),
         _fmt(sum(g["nameplate_mwh"] for g in by_block[index]), decimals=3),
         f"{row['pcs_count']} x {_fmt(row['pcs_kw'], ' kW', 0)}", _fmt(row["pcs_count"] * float(row["pcs_kw"]) / 1000)]
        for index, row in enumerate(rows, 1)
    ] + [["Total", count, len(plan["groups"]), _fmt(total, decimals=3), f"{pcs_total} PCS units", _fmt(pcs_mw)]],
        widths=[.12, .12, .13, .16, .31, .16], font_size=10, numeric_columns=(1, 2, 3, 5))
    topology = output.get("transformer_topology")
    doc.add_paragraph("Selected transformer interface: " + ("two windings with one common LV bus." if topology == "two_winding" else "three windings with two separate LV buses. PCS to winding allocation is a concept split pending OEM confirmation."))
    doc.add_paragraph("Transformer interfaces shown: one per actual AC Block. MVA rating, vector group, impedance, earthing, MV equipment and protection remain to be confirmed. Enclosure quantity and physical packaging are not established by this schedule.")
    doc.add_paragraph("Sections 7 and 8 include every actual AC Block. Each electrical connection concept is followed in the next section by the corresponding block arrangement, using the same saved group and PCS identifiers.", style="Caption")

    ordered_pages = sorted(diagram_pages, key=lambda page: page["ac_block_index"])
    for position, page in enumerate(ordered_pages):
        index = page["ac_block_index"]
        _diagram_section(doc, page["sld_svg"], brand, f"AC Block {index} electrical connection concept NOT FOR CONSTRUCTION",
                         chapter_title=REPORT_CHAPTERS[6] if position == 0 else None)
    for position, page in enumerate(ordered_pages):
        index = page["ac_block_index"]
        _diagram_section(doc, page["layout_svg"], brand, f"AC Block {index} diagrammatic cabinet arrangement NOT FOR CONSTRUCTION",
                         chapter_title=REPORT_CHAPTERS[7] if position == 0 else None)

    _portrait_section(doc, REPORT_CHAPTERS[8])
    doc.add_paragraph(f"The selected configuration comprises {len(rows)} AC Blocks and {count} physical cabinets. The block arrangements in Section 8 show saved equipment membership and relative zones only. No site envelope, land area, equipment separation or approved route is calculated for this cabinet alternative.")
    doc.add_paragraph("A dimensional site arrangement remains pending the verified site boundary, equipment footprints, access requirements and applicable project constraints. The symbolic interface reserves in Section 8 do not establish physical dimensions or separation distances.")
    _installation_basis(doc, data.get("installation_basis"))
    confirmations = [
        ("Equipment and electrical interfaces", "Confirm cabinet and PCS voltage/current limits, BMS coordination, transformer nameplate/vector group, earthing, switching and protection against the final OEM data."),
        ("Site integration and access", "Confirm the site boundary, equipment footprints, door and service sides, fire and maintenance access, cable routes, drainage and civil interfaces before any dimensional layout."),
        ("Performance and operating basis", "Verify net POI power, auxiliary demand and reactive capability. Agree dispatch, cycles, SOC window, temperature, augmentation and warranty terms separately from the saved model assessment."),
    ]
    for heading, explanation in confirmations:
        doc.add_heading(heading, level=2)
        doc.add_paragraph(explanation)
    doc.add_paragraph("Drawing status remains CONCEPT / NOT FOR CONSTRUCTION. No approved site plan, dimensional installation drawing or final electrical design is issued by this export.")
    _brand_headers(doc, brand)
    content = _doc_to_bytes(doc)
    assert_brand_clean_docx(content, brand)
    return content
