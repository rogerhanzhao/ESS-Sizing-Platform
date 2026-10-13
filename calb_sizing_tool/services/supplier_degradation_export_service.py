"""Presentation exports for the supplier review page using existing XLSX support."""
from collections import defaultdict
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Series, Reference
from openpyxl.chart.data_source import NumData, NumVal
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY = "163B4E"
TEAL = "007F82"
LIGHT = "EFF6F7"
GREY = "617481"
AMBER = "FFF2D8"
COLORS = ["007F82", "245BB2", "EA963E", "A65078", "6957AC", "7C9651"]
SERIES = {"lower_limit": "Lower limit", "mean": "Mean"}
PERCENT = "0.00%"


def _write(sheet, row, col, value, *, percent=False, bold=False, fill=None, color=NAVY):
    cell = sheet.cell(row, col, value)
    if isinstance(value, str):
        cell.data_type = "s"  # Source text is never an Excel formula.
    cell.font = Font(name="Calibri", size=10, color=color, bold=bold)
    cell.alignment = Alignment(horizontal="center" if col > 1 else "left", vertical="center", wrap_text=True)
    if percent and isinstance(value, (int, float)):
        cell.number_format = PERCENT
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    return cell


def _band(sheet, row, text, last=15, color=NAVY, height=28):
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last)
    _write(sheet, row, 1, text, bold=True, fill=color, color="FFFFFF")
    sheet.row_dimensions[row].height = height


def _note(sheet, row, text, last=15, height=28):
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last)
    _write(sheet, row, 1, text, color=GREY)
    sheet.row_dimensions[row].height = height


def _setup(sheet, *, last="O", landscape=True):
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins.left = sheet.page_margins.right = 0.25
    sheet.page_margins.top = sheet.page_margins.bottom = 0.3
    sheet.oddFooter.left.text = "CALB | Capacity retention reference"
    sheet.oddFooter.right.text = "Page &P of &N"
    sheet.oddFooter.left.size = sheet.oddFooter.right.size = 8
    sheet.sheet_properties.tabColor = TEAL
    for col in range(1, 20):
        from openpyxl.utils import get_column_letter
        sheet.column_dimensions[get_column_letter(col)].width = 11
    sheet.column_dimensions["A"].width = 12


def _common(curve):
    return curve.power_rate in (0.25, 0.33, 0.5) and curve.cycles_per_day <= 2


def _curve_sheet(book, name, curves, *, extended=False):
    sheet = book.create_sheet(name)
    _setup(sheet)
    sheet.freeze_panes = "B7"
    if extended:
        sheet.sheet_properties.tabColor = "C38B32"
    grouped = defaultdict(list)
    for curve in curves:
        grouped[curve.power_rate].append(curve)
    last_col = max(12, max(len(cases) for cases in grouped.values()) + 11)
    soc, basis = curves[0].soc_pct, curves[0].basis
    row = 1
    for section, (rate, cases) in enumerate(sorted(grouped.items())):
        if section:
            sheet.row_breaks.append(Break(id=row - 1))
        cases.sort(key=lambda c: c.cycles_per_day)
        _band(sheet, row, f"CALB | Cell capacity retention — {SERIES[basis]}", last=last_col)
        _note(sheet, row + 1, f"Source SOC label: {soc:.2f}%   |   Source Power Rate: {rate:g}   |   {'Extended operating cases' if extended else 'Common operating cases'}", last=last_col, height=24)
        _note(sheet, row + 2, "Reference data. SOC and Power Rate definitions require confirmation; Lower limit is not a confirmed warranty guarantee.", last=last_col, height=27)
        header = row + 4
        _write(sheet, header, 1, "Year", bold=True, fill=TEAL, color="FFFFFF")
        for col, c in enumerate(cases, 2):
            _write(sheet, header, col, f"{c.cycles_per_day:g} cycle{'s' if c.cycles_per_day != 1 else ''}/day", bold=True, fill=TEAL, color="FFFFFF")
        sheet.row_dimensions[header].height = 26
        years = sorted({y for c in cases for y in c.years})
        values = [dict(zip(c.years, c.retention_pct)) for c in cases]
        first = header + 1
        for offset, year in enumerate(years):
            r = first + offset
            sheet.row_dimensions[r].height = 14
            _write(sheet, r, 1, year, fill=LIGHT if offset % 2 == 0 else "FFFFFF")
            for col, source in enumerate(values, 2):
                v = source.get(year)
                cell = _write(sheet, r, col, "N/A" if v is None else v / 100, percent=True,
                    fill=AMBER if v is None else (LIGHT if offset % 2 == 0 else "FFFFFF"), color=GREY if v is None else NAVY)
                cell.border = Border(bottom=Side(style="hair", color="DBE5E8"))
        last = first + len(years) - 1
        chart = ScatterChart()
        chart.title = "Capacity retention"
        chart.x_axis.title = "Year"
        chart.y_axis.title = "Capacity retention (%)"
        chart.y_axis.numFmt = PERCENT
        chart.y_axis.scaling.min = 0
        chart.y_axis.scaling.max = 1
        chart.y_axis.majorUnit = 0.2
        chart.x_axis.scaling.min = min(years)
        chart.x_axis.scaling.max = max(years) if max(years) > 0 else 1
        chart.x_axis.majorUnit = 5 if max(years) >= 10 else 1
        chart.legend.position = "b"
        chart.style = 13
        chart.width, chart.height = 18, 11
        chart.scatterStyle = "line"
        chart.display_blanks = "gap"
        xref = Reference(sheet, min_col=1, min_row=first, max_row=last)
        for col, c in enumerate(cases, 2):
            series = Series(Reference(sheet, min_col=col, min_row=first, max_row=last), xref, title=f"{c.cycles_per_day:g} cycle{'s' if c.cycles_per_day != 1 else ''}/day")
            # Native Excel can render immediately, without recalculating refs.
            series.xVal.numRef.numCache = NumData(formatCode="0", ptCount=len(years),
                pt=[NumVal(idx=i, v=y) for i, y in enumerate(years)])
            source_values = dict(zip(c.years, c.retention_pct))
            series.yVal.numRef.numCache = NumData(formatCode=PERCENT, ptCount=len(years),
                pt=[NumVal(idx=i, v=source_values[y] / 100) for i, y in enumerate(years) if source_values.get(y) is not None])
            series.graphicalProperties.line.solidFill = COLORS[(col - 2) % len(COLORS)]
            series.graphicalProperties.line.width = 24000
            series.marker.symbol = "none"
            series.smooth = False
            chart.series.append(series)
        from openpyxl.utils import get_column_letter
        sheet.add_chart(chart, f"{get_column_letter(len(cases) + 3)}{header}")
        _note(sheet, last + 2, "N/A = no supplier value. No missing-year extension, interpolation or calendar/cycle recombination.", last=last_col, height=24)
        _note(sheet, last + 3, "Last source year: " + "  |  ".join(f"{c.cycles_per_day:g}/day: {c.last_supported_year if c.last_supported_year is not None else 'N/A'}" for c in cases), last=last_col, height=24)
        over = [c for c in cases if c.cycles_per_day * (2 / c.power_rate + 1) > 24]
        if over:
            _note(sheet, last + 4, "Conditional time check exceeds 24 h/day: " + "; ".join(f"{c.cycles_per_day:g}/day needs {c.cycles_per_day * (2 / c.power_rate + 1):.2f} h" for c in over) + ". Assumes 100% depth, equal rates and two 30-minute rests.", last=last_col, height=35)
        row = last + 6
    from openpyxl.utils import get_column_letter
    sheet.print_area = f"A1:{get_column_letter(last_col)}{row - 1}"
    return sheet


def _detail_sheet(book, title, headers, rows, *, percent_columns=(), widths=None):
    sheet = book.create_sheet(title)
    _setup(sheet)
    _band(sheet, 1, f"CALB | {title.replace('_', ' ')}", last=len(headers))
    _note(sheet, 2, "Source values for detailed review. Filter operating conditions as needed; no system guarantee is implied.", last=len(headers), height=30)
    for col, label in enumerate(headers, 1):
        _write(sheet, 4, col, label, bold=True, fill=TEAL, color="FFFFFF")
    sheet.row_dimensions[4].height = 30
    for r, values in enumerate(rows, 5):
        sheet.row_dimensions[r].height = max(20, min(80, ((len(str(values[-1])) // 40) + 1) * 14))
        for col, value in enumerate(values, 1):
            cell = _write(sheet, r, col, value, percent=col in percent_columns, fill=LIGHT if r % 2 else "FFFFFF")
            if isinstance(value, float) and col not in percent_columns:
                cell.number_format = "0.00"
    from openpyxl.utils import get_column_letter
    end = get_column_letter(len(headers))
    table = Table(displayName=title.replace("_", ""), ref=f"A4:{end}{len(rows) + 4}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    sheet.freeze_panes = "A5"
    for col, width in enumerate(widths or [17] * len(headers), 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.print_title_rows = "1:4"
    sheet.print_area = f"A1:{end}{len(rows) + 4}"
    return sheet


def export_supplier_review_workbook(dataset, curves=None, metadata=None):
    """All and selected exports share one formatted workbook contract."""
    curves = tuple(dataset.curves if curves is None else curves)
    if not curves:
        raise ValueError("Select at least one source curve.")
    book = Workbook()
    contents = book.active
    contents.title = "Start_here"
    _setup(contents, landscape=False)
    contents.page_setup.fitToHeight = 1
    _band(contents, 1, "CALB | Cell capacity retention", last=8, height=36)
    _note(contents, 2, "Supplier reference curves • Annual tables and comparison charts", last=8)
    metadata = metadata or {}
    _note(contents, 4, f"Version: {metadata.get('version_label') or 'Supplier reference'}", last=8)
    if metadata.get("filename"):
        _note(contents, 5, f"Source workbook: {metadata['filename']}", last=8)
    _note(contents, 7, f"Included: {len(curves)} source curves. Capacity retention and source SOC are shown as percentages with two decimal places.", last=8, height=34)
    _note(contents, 8, "SOC50 appears first. Each SOC / series has separate common and extended sheets; rates have separate printable sections.", last=8, height=34)
    _band(contents, 10, "Open a worksheet", last=8, color=TEAL)
    groups = defaultdict(list)
    for c in curves:
        groups[(c.soc_pct, c.basis, not _common(c))].append(c)
    ordered = sorted(groups, key=lambda k: (k[2], k[0] != 50, k[0], k[1] != "lower_limit"))
    for row, key in enumerate(ordered, 11):
        soc, basis, extended = key
        name = f"SOC{soc:g}_{'Lower' if basis == 'lower_limit' else 'Mean'}{'_Extended' if extended else ''}"
        _curve_sheet(book, name, groups[key], extended=extended)
        contents.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        cell = _write(contents, row, 1, f"SOC {soc:.2f}%  |  {SERIES[basis]}  |  {'Extended cases' if extended else 'Common cases'}  ({len(groups[key])} curves)", fill=LIGHT)
        cell.hyperlink = f"#'{name}'!A1"
        cell.font = Font(name="Calibri", size=11, color=TEAL, underline="single")
        contents.row_dimensions[row].height = 25
    row = len(ordered) + 13
    _band(contents, row, "Reading notes", last=8, color=TEAL)
    notes = ["Mean and Lower limit reproduce the supplier sheet names. Their statistical or warranty basis is not confirmed.",
             "SOC is a source label. Standby SOC versus cycle window is unconfirmed; SOC is not renamed DOD.",
             "Power Rate is a source label; nominal charging/discharging basis requires confirmation.",
             "N/A means no supplier value. Source checkpoints are retained without extrapolation or smoothing.",
             "Duty timing assumes 100% depth, equal nominal rates and two 30-minute rests per cycle; it is a conditional screen.",
             "No separate calendar-loss recombination or platform RTE replacement is made. Curves are reference data, not activated sizing data."]
    for offset, note in enumerate(notes, 1):
        _note(contents, row + offset, note, last=8, height=34)
    contents.print_area = f"A1:H{row + len(notes)}"
    contents.freeze_panes = "A11"
    annual_rows, duty_rows = [], []
    for c in sorted(curves, key=lambda c: (not _common(c), c.soc_pct != 50, c.soc_pct, c.basis != "lower_limit", c.power_rate, c.cycles_per_day)):
        for y, v, cell in zip(c.years, c.retention_pct, c.source_cells):
            annual_rows.append([c.soc_pct / 100, SERIES[c.basis], c.power_rate, c.cycles_per_day, y, "N/A" if v is None else v / 100, c.sheet.strip(), cell])
        hours = c.cycles_per_day * (2 / c.power_rate + 1)
        duty_rows.append([c.soc_pct / 100, SERIES[c.basis], c.power_rate, c.cycles_per_day,
            c.last_supported_year if c.last_supported_year is not None else "N/A", sum(v is None for v in c.retention_pct),
            hours, "Exceeds 24 h" if hours > 24 else "Fits assumed schedule", " | ".join(c.issues) or "—"])
    _detail_sheet(book, "Annual_data", ["Source SOC", "Source series", "Power Rate", "Cycles/day", "Year", "Capacity retention", "Source sheet", "Source cell"], annual_rows,
        percent_columns=(1, 6), widths=[16, 17, 15, 15, 10, 23, 26, 15])
    duty = _detail_sheet(book, "Duty_review", ["Source SOC", "Source series", "Power Rate", "Cycles/day", "Last source year", "Missing years", "Assumed hours/day", "Conditional time check", "Source observations"], duty_rows,
        percent_columns=(1,), widths=[15, 16, 14, 14, 19, 18, 22, 28, 42])
    for r in range(5, len(duty_rows) + 5):
        if duty.cell(r, 8).value == "Exceeds 24 h":
            duty.cell(r, 8).fill = PatternFill("solid", fgColor=AMBER)
    book.properties.title = "CALB cell capacity retention reference"
    book.properties.creator = "CALB Sizing Tool"
    book.properties.description = "Supplier source data formatted for review; not a warranty guarantee."
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()
